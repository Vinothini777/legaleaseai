import io
import os
import re
from pathlib import Path

import streamlit as st
from dotenv import load_dotenv
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from ai_core.gemini_generator import GeminiDocumentGenerator


# ============================================================
# PATH / ENVIRONMENT
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent

load_dotenv(PROJECT_ROOT / ".env")


# Load Gemini API key from Streamlit Secrets when deployed.
# Locally, .env can still provide GEMINI_API_KEY.
try:
    secret_key = st.secrets.get("GEMINI_API_KEY", "")
    if secret_key:
        os.environ["GEMINI_API_KEY"] = secret_key
except Exception:
    pass


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# CUSTOM CSS
# ============================================================

st.markdown(
    """
    <style>

    /* --------------------------------------------------------
       GLOBAL
    -------------------------------------------------------- */

    .stApp {
        background: #f4f6f9;
    }

    .main .block-container {
        max-width: 1180px;
        padding-top: 2rem;
        padding-bottom: 4rem;
        padding-left: 2rem;
        padding-right: 2rem;
    }

    /* --------------------------------------------------------
       SIDEBAR
    -------------------------------------------------------- */

    section[data-testid="stSidebar"] {
        background: #eef1f5;
        border-right: 1px solid #dfe3e8;
    }

    section[data-testid="stSidebar"] .block-container {
        padding-top: 2rem;
        padding-left: 1rem;
        padding-right: 1rem;
    }

    .sidebar-title {
        font-size: 1.15rem;
        font-weight: 700;
        color: #26364a;
        margin-bottom: 1.2rem;
    }

    .sidebar-note {
        margin-top: 2rem;
        padding-top: 1.2rem;
        border-top: 1px solid #cfd5dc;
        font-size: 0.78rem;
        line-height: 1.55;
        color: #6c7480;
    }

    /* --------------------------------------------------------
       HEADER
    -------------------------------------------------------- */

    .legal-header {
        text-align: center;
        margin-bottom: 2.2rem;
    }

    .legal-title {
        color: #26364a;
        font-size: 3rem;
        font-weight: 800;
        letter-spacing: -1px;
        margin-bottom: 0.25rem;
    }

    .legal-subtitle {
        color: #727b87;
        font-size: 1rem;
        margin-top: 0;
    }

    /* --------------------------------------------------------
       SECTION HEADINGS
    -------------------------------------------------------- */

    .form-section-title {
        color: #26364a;
        font-size: 1.15rem;
        font-weight: 750;
        margin-top: 0.4rem;
        margin-bottom: 0.15rem;
    }

    .form-helper {
        color: #7a838e;
        font-size: 0.78rem;
        margin-bottom: 0.45rem;
    }

    /* --------------------------------------------------------
       INPUTS
    -------------------------------------------------------- */

    div[data-testid="stTextInput"] input,
    div[data-testid="stTextArea"] textarea,
    div[data-baseweb="select"] {
        border-radius: 10px !important;
    }

    div[data-testid="stTextArea"] textarea {
        min-height: 150px;
        background: #ffffff;
        border: 1px solid #d8dde5;
        color: #26364a;
        font-size: 0.92rem;
        line-height: 1.55;
    }

    div[data-testid="stTextInput"] input {
        background: #ffffff;
        border: 1px solid #d8dde5;
        color: #26364a;
        min-height: 45px;
    }

    /* --------------------------------------------------------
       GENERATE BUTTON
    -------------------------------------------------------- */

    div.stButton > button {
        width: 100%;
        min-height: 48px;
        border-radius: 9px;
        border: none;
        background: #ff4f52;
        color: white;
        font-weight: 700;
        font-size: 0.98rem;
        transition: all 0.2s ease;
    }

    div.stButton > button:hover {
        background: #e94346;
        color: white;
        border: none;
    }

    div.stButton > button:focus {
        color: white;
        border: none;
        box-shadow: none;
    }

    /* --------------------------------------------------------
       GENERATED DOCUMENT
    -------------------------------------------------------- */

    .generated-card {
        background: #ffffff;
        border: 1px solid #e0e4ea;
        border-radius: 12px;
        padding: 1.6rem;
        margin-top: 1.2rem;
        box-shadow: 0 2px 10px rgba(30, 40, 55, 0.04);
    }

    .generated-text {
        color: #26364a;
        font-size: 0.94rem;
        line-height: 1.75;
        white-space: pre-wrap;
    }

    .document-title {
        color: #26364a;
        font-size: 1.4rem;
        font-weight: 800;
        margin-bottom: 0.8rem;
    }

    /* --------------------------------------------------------
       DOWNLOAD BUTTONS
    -------------------------------------------------------- */

    div[data-testid="stDownloadButton"] button {
        width: 100%;
        border-radius: 9px;
        min-height: 45px;
        font-weight: 650;
    }

    /* --------------------------------------------------------
       STATUS
    -------------------------------------------------------- */

    .gemini-status {
        background: #e9f8ee;
        border: 1px solid #bde8ca;
        color: #18733b;
        padding: 0.85rem 1rem;
        border-radius: 9px;
        margin-top: 1rem;
        margin-bottom: 1rem;
        font-size: 0.9rem;
    }

    .demo-status {
        background: #fff8dc;
        border: 1px solid #f0df9b;
        color: #765f16;
        padding: 0.85rem 1rem;
        border-radius: 9px;
        margin-top: 1rem;
        margin-bottom: 1rem;
        font-size: 0.9rem;
    }

    /* --------------------------------------------------------
       DISCLAIMER
    -------------------------------------------------------- */

    .legal-disclaimer {
        color: #7a838e;
        font-size: 0.78rem;
        line-height: 1.55;
        padding: 1rem 0;
        border-top: 1px solid #e0e4e9;
        margin-top: 2rem;
    }

    /* --------------------------------------------------------
       MOBILE
    -------------------------------------------------------- */

    @media (max-width: 768px) {

        .main .block-container {
            padding-left: 1rem;
            padding-right: 1rem;
        }

        .legal-title {
            font-size: 2.2rem;
        }

    }

    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# SESSION STATE
# ============================================================

if "generated" not in st.session_state:
    st.session_state.generated = False

if "document" not in st.session_state:
    st.session_state.document = ""

if "document_type" not in st.session_state:
    st.session_state.document_type = ""

if "generation_mode" not in st.session_state:
    st.session_state.generation_mode = ""


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.markdown(
        '<div class="sidebar-title">Document Setup</div>',
        unsafe_allow_html=True,
    )

    document_types = [
        "Rental Agreement",
        "Employment Agreement",
        "Non-Disclosure Agreement",
        "Service Agreement",
        "Lease Agreement",
        "Partnership Agreement",
        "Loan Agreement",
        "Sale Agreement",
        "Freelance Agreement",
        "Affidavit",
        "Power of Attorney",
        "Other Legal Document",
    ]

    document_type = st.selectbox(
        "Document type",
        document_types,
        index=0,
    )

    jurisdiction = st.text_input(
        "Jurisdiction",
        value="India",
    )

    st.markdown(
        """
        <div class="sidebar-note">
        LegalEase creates AI-generated drafts for informational purposes only.
        Important documents should be reviewed by a qualified legal professional.
        </div>
        """,
        unsafe_allow_html=True,
    )


# ============================================================
# HEADER
# ============================================================

st.markdown(
    """
    <div class="legal-header">
        <div class="legal-title">⚖️ LegalEase</div>
        <div class="legal-subtitle">
            AI-powered legal document drafting assistant
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# FORM
# ============================================================

left_column, right_column = st.columns(
    [1, 1],
    gap="large",
)


# ============================================================
# LEFT COLUMN
# ============================================================

with left_column:

    st.markdown(
        '<div class="form-section-title">1. Parties</div>',
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="form-helper">Enter the parties involved</div>',
        unsafe_allow_html=True,
    )

    parties = st.text_area(
        "Parties",
        placeholder=(
            "Example:\n"
            "Landlord: Priya Kumar\n"
            "Tenant: Arjun Mehta"
        ),
        height=150,
        label_visibility="collapsed",
    )

    st.markdown(
        '<div class="form-section-title">2. Terms</div>',
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="form-helper">Enter the terms and conditions</div>',
        unsafe_allow_html=True,
    )

    terms = st.text_area(
        "Terms",
        placeholder=(
            "Describe salary, duties, payment terms, confidentiality, "
            "termination, etc."
        ),
        height=260,
        label_visibility="collapsed",
    )


# ============================================================
# RIGHT COLUMN
# ============================================================

with right_column:

    st.markdown(
        '<div class="form-section-title">3. Effective Date</div>',
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="form-helper">Effective date</div>',
        unsafe_allow_html=True,
    )

    effective_date = st.text_input(
        "Effective Date",
        value="1 October 2026",
        label_visibility="collapsed",
    )

    st.markdown(
        '<div class="form-section-title">4. Additional Instructions</div>',
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="form-helper">Optional instructions</div>',
        unsafe_allow_html=True,
    )

    additional_instructions = st.text_area(
        "Additional Instructions",
        placeholder=(
            "Example: Use formal language and numbered clauses."
        ),
        height=260,
        label_visibility="collapsed",
    )


# ============================================================
# GENERATE BUTTON
# ============================================================

st.write("")

generate_clicked = st.button(
    "⚖ Generate Legal Document",
    use_container_width=True,
)


# ============================================================
# GEMINI GENERATION
# ============================================================

if generate_clicked:

    # Basic validation
    if not parties.strip():
        st.error("Please enter the parties involved.")

    elif not terms.strip():
        st.error("Please enter the terms and conditions.")

    elif not effective_date.strip():
        st.error("Please enter the effective date.")

    elif not jurisdiction.strip():
        st.error("Please enter the jurisdiction.")

    else:

        with st.spinner("Generating your legal document..."):

            try:

                # Get Gemini API key from Streamlit Secrets.
                # Fall back to .env locally.
                try:
                    gemini_key = st.secrets.get(
                        "GEMINI_API_KEY",
                        "",
                    )
                except Exception:
                    gemini_key = ""

                if not gemini_key:
                    gemini_key = os.getenv(
                        "GEMINI_API_KEY",
                        "",
                    )

                generator = GeminiDocumentGenerator(
                    api_key=gemini_key,
                )

                generated_document = generator.generate_document(
                    document_type=document_type,
                    parties=parties,
                    terms=terms,
                    effective_date=effective_date,
                    jurisdiction=jurisdiction,
                    additional_instructions=additional_instructions,
                )

                # Store generated result
                st.session_state.document = generated_document
                st.session_state.document_type = document_type
                st.session_state.generation_mode = generator.mode
                st.session_state.generated = True

                st.success(
                    "Document generated successfully!"
                )

            except Exception as exc:

                st.session_state.generated = False
                st.session_state.document = ""
                st.session_state.generation_mode = ""

                st.error(
                    "Gemini could not generate the document."
                )

                st.exception(exc)


# ============================================================
# GENERATED DOCUMENT
# ============================================================

if st.session_state.generated and st.session_state.document:

    st.markdown(
        "---"
    )

    st.markdown(
        """
        <div class="form-section-title"
             style="font-size:1.45rem; margin-bottom:1rem;">
            Generated Document
        </div>
        """,
        unsafe_allow_html=True,
    )

    # --------------------------------------------------------
    # Generation status
    # --------------------------------------------------------

    mode = st.session_state.generation_mode

    if mode == "gemini":

        st.markdown(
            """
            <div class="gemini-status">
                ✓ Document generated using Gemini AI.
            </div>
            """,
            unsafe_allow_html=True,
        )

    elif mode == "demo":

        st.markdown(
            """
            <div class="demo-status">
                Demo mode is active. Gemini AI was not used.
            </div>
            """,
            unsafe_allow_html=True,
        )

    # --------------------------------------------------------
    # Document display
    # --------------------------------------------------------

    document_text = st.session_state.document

    st.markdown(
        '<div class="generated-card">',
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="document-title">Legal Document</div>',
        unsafe_allow_html=True,
    )

    st.markdown(
        f'<div class="generated-text">{document_text}</div>',
        unsafe_allow_html=True,
    )

    st.markdown(
        '</div>',
        unsafe_allow_html=True,
    )


# ============================================================
# DOCX GENERATION
# ============================================================

def create_docx(
    text: str,
    document_type: str,
) -> bytes:

    document = Document()

    # Margins
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # Normal font
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(11)

    # Title
    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(document_type.upper())
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)

    document.add_paragraph("")

    # Content
    for line in text.splitlines():

        cleaned = line.strip()

        if not cleaned:
            document.add_paragraph("")
            continue

        paragraph = document.add_paragraph()

        # Detect headings
        if (
            re.match(r"^\d+[\.\)]\s+", cleaned)
            or cleaned.upper() == cleaned
            or cleaned.endswith(":")
        ):
            run = paragraph.add_run(cleaned)
            run.bold = True
        else:
            run = paragraph.add_run(cleaned)

        run.font.name = "Arial"
        run.font.size = Pt(11)

    # Save to memory
    output = io.BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


# ============================================================
# PDF FONT
# ============================================================

def register_pdf_font():

    possible_fonts = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/ARIAL.TTF"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
    ]

    for font_path in possible_fonts:

        if font_path.exists():

            try:

                pdfmetrics.registerFont(
                    TTFont(
                        "LegalEaseFont",
                        str(font_path),
                    )
                )

                return "LegalEaseFont"

            except Exception:
                continue

    return "Helvetica"


# ============================================================
# PDF GENERATION
# ============================================================

def create_pdf(
    text: str,
    document_type: str,
) -> bytes:

    output = io.BytesIO()

    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
    )

    font_name = register_pdf_font()

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "LegalEaseTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=16,
        leading=20,
        alignment=TA_CENTER,
        spaceAfter=20,
    )

    body_style = ParagraphStyle(
        "LegalEaseBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=10.5,
        leading=16,
        spaceAfter=8,
    )

    heading_style = ParagraphStyle(
        "LegalEaseHeading",
        parent=body_style,
        fontName=font_name,
        fontSize=11,
        leading=16,
        spaceBefore=8,
        spaceAfter=7,
    )

    story = []

    story.append(
        Paragraph(
            document_type.upper(),
            title_style,
        )
    )

    for raw_line in text.splitlines():

        line = raw_line.strip()

        if not line:
            story.append(
                Spacer(1, 7)
            )
            continue

        # Escape HTML-sensitive characters
        safe_line = (
            line
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        if (
            re.match(r"^\d+[\.\)]\s+", line)
            or line.upper() == line
            or line.endswith(":")
        ):

            story.append(
                Paragraph(
                    f"<b>{safe_line}</b>",
                    heading_style,
                )
            )

        else:

            story.append(
                Paragraph(
                    safe_line,
                    body_style,
                )
            )

    pdf.build(story)

    output.seek(0)

    return output.getvalue()


# ============================================================
# DOWNLOAD SECTION
# ============================================================

if st.session_state.generated and st.session_state.document:

    st.markdown(
        ""
    )

    col1, col2 = st.columns(
        2,
        gap="medium",
    )

    safe_name = re.sub(
        r"[^a-zA-Z0-9_-]+",
        "_",
        st.session_state.document_type.strip(),
    )

    if not safe_name:
        safe_name = "legal_document"

    # DOCX
    with col1:

        try:

            docx_bytes = create_docx(
                st.session_state.document,
                st.session_state.document_type,
            )

            st.download_button(
                label="⬇ Download DOCX",
                data=docx_bytes,
                file_name=f"{safe_name}.docx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                use_container_width=True,
            )

        except Exception as exc:

            st.error(
                f"Could not create DOCX: {exc}"
            )

    # PDF
    with col2:

        try:

            pdf_bytes = create_pdf(
                st.session_state.document,
                st.session_state.document_type,
            )

            st.download_button(
                label="⬇ Download PDF",
                data=pdf_bytes,
                file_name=f"{safe_name}.pdf",
                mime="application/pdf",
                use_container_width=True,
            )

        except Exception as exc:

            st.error(
                f"Could not create PDF: {exc}"
            )


# ============================================================
# FOOTER
# ============================================================

st.markdown(
    """
    <div class="legal-disclaimer">
        LegalEase creates AI-generated drafts for informational purposes only.
        This application does not provide legal advice. Important legal
        documents should be reviewed by a qualified legal professional
        before signing or relying upon them.
    </div>
    """,
    unsafe_allow_html=True,
)