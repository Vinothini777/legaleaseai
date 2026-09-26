from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib import colors


def format_txt(text: str) -> bytes:
    """
    Convert document text into a UTF-8 TXT file.
    """
    return text.encode("utf-8")


def _is_heading(line: str) -> bool:
    """
    Detect common legal-document headings.
    """
    stripped = line.strip()

    if not stripped:
        return False

    if re.match(r"^(SECTION|ARTICLE)\s+\d+", stripped, re.IGNORECASE):
        return True

    if re.match(r"^\d+\.\s+[A-Z]", stripped):
        return True

    if stripped.isupper() and len(stripped) < 120:
        return True

    return False


def _clean_text(text: str) -> str:
    """
    Normalize common Unicode characters.
    """
    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " ",
        "\ufeff": "",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text.strip()


def format_docx(
    text: str,
    document_type: str = "Legal Document",
    logo_path=None,
    terms: str = "",
) -> bytes:
    """
    Generate a professionally formatted DOCX document.
    """

    document = Document()

    # -------------------------------------------------
    # Page margins
    # -------------------------------------------------
    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # -------------------------------------------------
    # Normal style
    # -------------------------------------------------
    normal = document.styles["Normal"]

    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    # -------------------------------------------------
    # Optional logo
    # -------------------------------------------------
    if logo_path:

        logo = Path(logo_path)

        if logo.exists():

            paragraph = document.add_paragraph()

            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = paragraph.add_run()

            run.add_picture(
                str(logo),
                width=Inches(1.4),
            )

    # -------------------------------------------------
    # Title
    # -------------------------------------------------
    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        _clean_text(document_type).upper()
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    # -------------------------------------------------
    # Document content
    # -------------------------------------------------
    cleaned = _clean_text(text)

    paragraphs = cleaned.split("\n")

    for line in paragraphs:

        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15

        if _is_heading(line):

            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

            run = paragraph.add_run(line)

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        else:

            run = paragraph.add_run(line)

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    # -------------------------------------------------
    # Terms table
    # -------------------------------------------------
    if terms and terms.strip():

        document.add_paragraph()

        heading = document.add_paragraph()

        run = heading.add_run(
            "KEY TERMS"
        )

        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

        term_items = [
            item.strip()
            for item in terms.split(";")
            if item.strip()
        ]

        if term_items:

            table = document.add_table(
                rows=1,
                cols=2,
            )

            table.style = "Table Grid"

            table.rows[0].cells[0].text = "No."
            table.rows[0].cells[1].text = "Term"

            for index, term in enumerate(
                term_items,
                start=1,
            ):

                cells = table.add_row().cells

                cells[0].text = str(index)
                cells[1].text = term

    # -------------------------------------------------
    # Footer
    # -------------------------------------------------
    footer = section.footer

    footer_paragraph = footer.paragraphs[0]

    footer_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer_paragraph.add_run(
        "LegalEase - Draft for review"
    )

    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(8)

    # -------------------------------------------------
    # Save to memory
    # -------------------------------------------------
    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


def format_pdf(
    text: str,
    document_type: str = "Legal Document",
    logo_path=None,
    terms: str = "",
) -> bytes:
    """
    Generate a professionally formatted PDF document.
    """

    output = BytesIO()

    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=document_type,
        author="LegalEase",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "LegalEaseTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        spaceAfter=15,
    )

    heading_style = ParagraphStyle(
        "LegalEaseHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        spaceBefore=8,
        spaceAfter=6,
    )

    body_style = ParagraphStyle(
        "LegalEaseBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        spaceAfter=7,
    )

    story = []

    # -------------------------------------------------
    # Title
    # -------------------------------------------------
    story.append(
        Paragraph(
            _clean_text(document_type).upper(),
            title_style,
        )
    )

    # -------------------------------------------------
    # Content
    # -------------------------------------------------
    cleaned = _clean_text(text)

    paragraphs = cleaned.split("\n")

    for line in paragraphs:

        line = line.strip()

        if not line:
            story.append(Spacer(1, 5))
            continue

        # Escape PDF-sensitive characters
        safe_line = (
            line.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        if _is_heading(line):

            story.append(
                Paragraph(
                    safe_line,
                    heading_style,
                )
            )

        else:

            story.append(
                Paragraph(
                    safe_line,
                    body_style,
                )
            )

    # -------------------------------------------------
    # Key terms table
    # -------------------------------------------------
    if terms and terms.strip():

        story.append(
            Spacer(1, 10)
        )

        story.append(
            Paragraph(
                "KEY TERMS",
                heading_style,
            )
        )

        term_items = [
            item.strip()
            for item in terms.split(";")
            if item.strip()
        ]

        if term_items:

            data = [
                ["No.", "Term"]
            ]

            for index, term in enumerate(
                term_items,
                start=1,
            ):

                safe_term = (
                    term.replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;")
                )

                data.append(
                    [
                        str(index),
                        Paragraph(
                            safe_term,
                            body_style,
                        ),
                    ]
                )

            table = Table(
                data,
                colWidths=[
                    15 * mm,
                    145 * mm,
                ],
                repeatRows=1,
            )

            table.setStyle(
                TableStyle(
                    [
                        (
                            "GRID",
                            (0, 0),
                            (-1, -1),
                            0.5,
                            colors.grey,
                        ),
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.lightgrey,
                        ),
                        (
                            "FONTNAME",
                            (0, 0),
                            (-1, 0),
                            "Helvetica-Bold",
                        ),
                        (
                            "FONTNAME",
                            (0, 1),
                            (0, -1),
                            "Helvetica",
                        ),
                        (
                            "VALIGN",
                            (0, 0),
                            (-1, -1),
                            "TOP",
                        ),
                        (
                            "LEFTPADDING",
                            (0, 0),
                            (-1, -1),
                            5,
                        ),
                        (
                            "RIGHTPADDING",
                            (0, 0),
                            (-1, -1),
                            5,
                        ),
                        (
                            "TOPPADDING",
                            (0, 0),
                            (-1, -1),
                            5,
                        ),
                        (
                            "BOTTOMPADDING",
                            (0, 0),
                            (-1, -1),
                            5,
                        ),
                    ]
                )
            )

            story.append(table)

    # -------------------------------------------------
    # Footer callback
    # -------------------------------------------------
    def add_footer(canvas, doc):
        canvas.saveState()

        canvas.setFont(
            "Helvetica",
            8,
        )

        canvas.drawCentredString(
            A4[0] / 2,
            10 * mm,
            "LegalEase - Draft for review",
        )

        canvas.restoreState()

    pdf.build(
        story,
        onFirstPage=add_footer,
        onLaterPages=add_footer,
    )

    output.seek(0)

    return output.getvalue()